import openpyxl

wb = openpyxl.Workbook()
sheet = wb.active
for i in range(1, 11):
    sheet['A' + str(i)] = i

ref_object = openpyxl.chart.Reference(sheet, min_col=1, min_row=1, max_col=1,
max_row=10)
series_object = openpyxl.chart.Series(ref_object, title='First Series')

chart_object = openpyxl.chart.BarChart()
chart_object.title = 'My Chart'
chart_object.append(series_object)

for i in range(1, 401):
    sheet['B' + str(i)] = i**2
ref_object_1 = openpyxl.chart.Reference(sheet, min_col=2, min_row=1, max_col=2, max_row=400)
series_object_1 = openpyxl.chart.Series(ref_object_1, title='Pizza Test')
chart_object_1 = openpyxl.chart.LineChart()
chart_object_1.title = 'My Pizza Test'
chart_object_1.append(series_object_1)
sheet.add_chart(chart_object_1, 'C20')
sheet.add_chart(chart_object, 'C5')
wb.save('sample_chart.xlsx')